"""Deterministic quotation rendering: template copy -> fill -> save -> verify.

No AI here on purpose: numbers and customer data on a quotation must be exact and reproducible.
"""
from __future__ import annotations

import datetime as dt
import os
import re
import shutil
import tempfile

import jinja2
from docx import Document
from docxtpl import DocxTemplate

from common.quote_logic import format_qty, format_vnd, validate_and_compute, vnd_in_words

TEMPLATE_PATH = os.environ.get("TEMPLATE_PATH", "templates/quotation_template.docx")


class PermanentError(Exception):
    """Retrying will not help (bad data, broken template)."""


def build_context(payload: dict) -> dict:
    # Defense in depth: re-validate and recompute on the agent. If the numbers the web app sent
    # disagree with our own calculation, refuse to produce a document.
    recomputed = validate_and_compute(payload)
    for key in ("subtotal", "vat_amount", "total"):
        if recomputed[key] != payload[key]:
            raise PermanentError(f"Số liệu không khớp ({key}: {payload[key]} != {recomputed[key]})")
    quote_date = dt.datetime.strptime(payload["quote_date"], "%d/%m/%Y")
    items = [{**it, "quantity_fmt": format_qty(it["quantity"]), "unit_price_fmt": format_vnd(it["unit_price"]), "amount_fmt": format_vnd(it["amount"])}
             for it in recomputed["items"]]
    return {
        **payload,
        "items": items,
        "subtotal_fmt": format_vnd(payload["subtotal"]),
        "vat_amount_fmt": format_vnd(payload["vat_amount"]),
        "total_fmt": format_vnd(payload["total"]),
        "total_words": vnd_in_words(payload["total"]),
        "valid_until": (quote_date + dt.timedelta(days=payload["validity_days"])).strftime("%d/%m/%Y"),
    }


def docx_text(path: str) -> str:
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def verify_output(path: str, ctx: dict) -> None:
    """Check the produced file really is a complete quotation before we report success."""
    if not os.path.exists(path) or os.path.getsize(path) == 0:
        raise RuntimeError("File đầu ra không tồn tại hoặc rỗng")
    try:
        text = docx_text(path)
    except Exception as e:  # corrupt zip / xml
        raise RuntimeError(f"File đầu ra không mở được: {e}") from e
    if re.search(r"\{\{|\{%|%\}|\}\}", text):
        raise PermanentError("Còn placeholder chưa được thay thế trong file")
    must_contain = [ctx["quote_no"], ctx["customer"]["name"], ctx["total_fmt"], ctx["total_words"]]
    for it in ctx["items"]:
        must_contain += [it["name"], it["quantity_fmt"], it["amount_fmt"]]
    missing = [s for s in must_contain if s not in text]
    if missing:
        raise PermanentError(f"File thiếu nội dung bắt buộc: {missing}")


def render_quote(payload: dict, out_dir: str) -> tuple[str, dict]:
    """Returns (path_to_docx, context). Works on a temp copy, then moves into place atomically."""
    if not os.path.exists(TEMPLATE_PATH):
        raise PermanentError(f"Không tìm thấy template {TEMPLATE_PATH}")
    ctx = build_context(payload)
    os.makedirs(out_dir, exist_ok=True)
    final_path = os.path.join(out_dir, f"{payload['quote_no']}.docx")
    with tempfile.TemporaryDirectory(dir=out_dir) as work:
        working_copy = os.path.join(work, "working.docx")
        shutil.copyfile(TEMPLATE_PATH, working_copy)  # never touch the master template
        tpl = DocxTemplate(working_copy)
        try:
            tpl.render(ctx, jinja_env=_jinja_env(), autoescape=True)
        except Exception as e:
            raise PermanentError(f"Lỗi điền template: {e}") from e
        rendered = os.path.join(work, "rendered.docx")
        tpl.save(rendered)
        verify_output(rendered, ctx)
        os.replace(rendered, final_path)
    return final_path, ctx


def _jinja_env():
    # StrictUndefined: a typo in a template placeholder fails loudly instead of printing blank.
    return jinja2.Environment(undefined=jinja2.StrictUndefined, autoescape=True)
