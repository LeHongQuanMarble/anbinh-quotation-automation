"""Build the (fictional) company quotation template: templates/quotation_template.docx.

In real life the company already owns a Word template; the only work needed is to put
Jinja placeholders ({{ ... }}) where the data goes. This script exists so the repo is
self-contained and the template is reproducible.
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BRAND = RGBColor(0x0B, 0x4F, 0x8A)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def para(doc_or_cell, text="", bold=False, size=10.5, align=None, color=None, space_after=2):
    p = doc_or_cell.add_paragraph()
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def main(path="templates/quotation_template.docx"):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(10.5)

    para(doc, "CÔNG TY CỔ PHẦN AN BÌNH CHEMTECH (DEMO)", bold=True, size=13, color=BRAND)
    para(doc, "Địa chỉ: 123 Đường Hóa Chất, KCN Demo, TP. Hồ Chí Minh · MST: 0000000000 (giả lập)", size=9)
    para(doc, "Điện thoại: 028 0000 0000 · Email: sales@anbinh-demo.example", size=9, space_after=10)

    para(doc, "BẢNG BÁO GIÁ", bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, color=BRAND)
    para(doc, "Số: {{ quote_no }}    Ngày: {{ quote_date }}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    info = doc.add_table(rows=0, cols=2)
    for label, value in [
        ("Kính gửi:", "{{ customer.name }}"),
        ("Mã số thuế:", "{{ customer.tax_code }}"),
        ("Địa chỉ:", "{{ customer.address }}"),
        ("Người liên hệ:", "{{ customer.contact_name }} - {{ customer.phone }} - {{ customer.email }}"),
    ]:
        row = info.add_row().cells
        row[0].text, row[1].text = label, value
        row[0].paragraphs[0].runs[0].bold = True
        row[0].width, row[1].width = Cm(3.5), Cm(13.9)
    para(doc, "")
    para(doc, "Công ty An Bình Chemtech trân trọng gửi đến Quý khách bảng báo giá như sau:", space_after=6)

    headers = ["STT", "Tên hàng", "Quy cách", "ĐVT", "Số lượng", "Đơn giá (VND)", "Thành tiền (VND)"]
    widths = [Cm(1.1), Cm(4.2), Cm(4.4), Cm(1.3), Cm(1.8), Cm(2.3), Cm(2.6)]
    t = doc.add_table(rows=4, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(c, "DCE8F4")
    # docxtpl row loop: the {%tr %} rows are removed at render time, the middle row is repeated.
    t.rows[1].cells[0].text = "{%tr for it in items %}"
    data = ["{{ it.no }}", "{{ it.name }}", "{{ it.spec }}", "{{ it.unit }}", "{{ it.quantity_fmt }}",
            "{{ it.unit_price_fmt }}", "{{ it.amount_fmt }}"]
    for i, v in enumerate(data):
        t.rows[2].cells[i].text = v
        if i >= 4:
            t.rows[2].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    t.rows[3].cells[0].text = "{%tr endfor %}"
    for label, value, bold in [("Cộng tiền hàng", "{{ subtotal_fmt }}", False),
                               ("Thuế GTGT {{ vat_rate }}%", "{{ vat_amount_fmt }}", False),
                               ("TỔNG CỘNG THANH TOÁN", "{{ total_fmt }}", True)]:
        cells = t.add_row().cells
        merged = cells[0].merge(cells[5])
        merged.text = label
        merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        merged.paragraphs[0].runs[0].bold = bold
        cells[6].text = value
        cells[6].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cells[6].paragraphs[0].runs[0].bold = bold
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

    para(doc, "")
    p = para(doc, "Bằng chữ: ", bold=True)
    r = p.add_run("{{ total_words }}.")
    r.italic = True

    para(doc, "ĐIỀU KIỆN THƯƠNG MẠI", bold=True, color=BRAND, space_after=4)
    for label, value in [("Thanh toán", "{{ payment_terms }}"), ("Giao hàng", "{{ delivery_terms }}"),
                         ("Thời gian giao hàng", "{{ delivery_time }}"),
                         ("Hiệu lực báo giá", "{{ validity_days }} ngày kể từ ngày {{ quote_date }} (đến {{ valid_until }})")]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)
    para(doc, "{% if notes %}Ghi chú: {{ notes }}{% endif %}", space_after=14)

    sign = doc.add_table(rows=1, cols=2)
    left, right = sign.rows[0].cells
    left.text = "XÁC NHẬN CỦA KHÁCH HÀNG\n(Ký, ghi rõ họ tên)"
    right.text = "ĐẠI DIỆN AN BÌNH CHEMTECH\n\n\n\n{{ sales.name }}\nNhân viên kinh doanh"
    for c in (left, right):
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].runs[0].bold = True

    doc.save(path)
    print("wrote", path)


if __name__ == "__main__":
    main()
